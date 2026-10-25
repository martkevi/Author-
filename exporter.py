"""
Manuscript exporter for Telegram Book Studio.

Supported formats:
- TXT
- DOCX
- EPUB
- PDF

The exporter retrieves saved chapters from SQLite and assembles
them into a manuscript in chapter order.

Exporting a manuscript does not automatically publish it on KDP.
The generated file must still be reviewed and prepared for the
chosen publishing format.
"""

import re
from pathlib import Path

import config
import storage


SUPPORTED_FORMATS = {"txt", "docx", "epub", "pdf"}


class ExportError(Exception):
    """Raised when a manuscript cannot be exported."""


def _safe_filename(title: str) -> str:
    """Convert a book title into a safe filename."""
    filename = re.sub(r"[^\w\s-]", "", title, flags=re.UNICODE)
    filename = re.sub(r"[\s-]+", "_", filename.strip())
    filename = filename.strip("_")

    return (filename[:100] or "untitled_book")


def _get_manuscript(
    book_id: int,
    user_id: int,
) -> tuple[dict, list[dict]]:
    """Retrieve a user's book and its saved chapters."""
    book = storage.get_book(book_id, user_id)

    if book is None:
        raise ExportError(
            "Book not found or you do not have permission to access it."
        )

    chapters = storage.list_chapters(book_id, user_id)

    if chapters is None:
        raise ExportError(
            "Could not retrieve the book's chapters."
        )

    if not chapters:
        raise ExportError(
            "This book has no saved chapters to export."
        )

    empty_chapters = [
        chapter["chapter_number"]
        for chapter in chapters
        if not chapter["content"].strip()
    ]

    if empty_chapters:
        raise ExportError(
            "Some chapters have no content: "
            + ", ".join(map(str, empty_chapters))
        )

    return book, chapters


def _build_plain_text(
    book: dict,
    chapters: list[dict],
) -> str:
    """Assemble the book as readable plain text."""
    parts = [
        book["title"].strip(),
        "",
        f"Genre: {book['genre']}",
        f"Target audience: {book['audience']}",
        "",
    ]

    if book.get("premise"):
        parts.extend(
            [
                "BOOK DESCRIPTION",
                "",
                book["premise"].strip(),
                "",
            ]
        )

    parts.extend(
        [
            "=" * 60,
            "",
        ]
    )

    for chapter in chapters:
        parts.extend(
            [
                chapter["title"].strip()
                or f"Chapter {chapter['chapter_number']}",
                "",
                chapter["content"].strip(),
                "",
                "",
            ]
        )

    return "\n".join(parts).rstrip() + "\n"


def _export_txt(
    output_path: Path,
    book: dict,
    chapters: list[dict],
) -> None:
    """Write a plain-text manuscript."""
    text = _build_plain_text(book, chapters)

    output_path.write_text(
        text,
        encoding="utf-8",
    )


def _export_docx(
    output_path: Path,
    book: dict,
    chapters: list[dict],
) -> None:
    """Create an editable Microsoft Word manuscript."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt

    except ImportError as exc:
        raise ExportError(
            "DOCX export requires python-docx. "
            "Install it using requirements.txt."
        ) from exc

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Georgia"
    normal_style.font.size = Pt(11)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_paragraph.add_run(book["title"])
    title_run.bold = True
    title_run.font.size = Pt(24)

    document.add_paragraph()

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run(f"Genre: {book['genre']}\n")
    metadata.add_run(f"Audience: {book['audience']}")

    if book.get("premise"):
        document.add_page_break()
        document.add_heading("Book Description", level=1)
        document.add_paragraph(book["premise"].strip())

    for chapter in chapters:
        document.add_page_break()

        heading = document.add_heading(
            chapter["title"].strip()
            or f"Chapter {chapter['chapter_number']}",
            level=1,
        )

        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for paragraph_text in chapter["content"].split("\n"):
            paragraph_text = paragraph_text.strip()

            if paragraph_text:
                paragraph = document.add_paragraph(
                    paragraph_text
                )
                paragraph.paragraph_format.space_after = Pt(8)

    document.save(str(output_path))


def _export_epub(
    output_path: Path,
    book: dict,
    chapters: list[dict],
) -> None:
    """Create a basic EPUB 3-compatible ebook."""
    try:
        from ebooklib import epub

    except ImportError as exc:
        raise ExportError(
            "EPUB export requires EbookLib. "
            "Install it using requirements.txt."
        ) from exc

    from html import escape

    epub_book = epub.EpubBook()

    epub_book.set_identifier(
        f"telegram-book-studio-{book['id']}"
    )

    epub_book.set_title(book["title"])
    epub_book.set_language("en")
    epub_book.add_author("Author")

    intro = epub.EpubHtml(
        title="Book Description",
        file_name="description.xhtml",
        lang="en",
    )

    description = escape(book.get("premise") or "")

    intro.content = (
        "<html>"
        "<head><title>Book Description</title></head>"
        "<body>"
        f"<h1>{escape(book['title'])}</h1>"
        f"<p>{description}</p>"
        "</body>"
        "</html>"
    )

    epub_book.add_item(intro)

    spine = ["nav", intro]
    chapter_items = []

    for chapter in chapters:
        number = chapter["chapter_number"]

        chapter_title = (
            chapter["title"].strip()
            or f"Chapter {number}"
        )

        paragraph_html = "".join(
            f"<p>{escape(paragraph.strip())}</p>"
            for paragraph in chapter["content"].split("\n")
            if paragraph.strip()
        )

        item = epub.EpubHtml(
            title=chapter_title,
            file_name=f"chapter_{number}.xhtml",
            lang="en",
        )

        item.content = (
            "<html>"
            "<head>"
            f"<title>{escape(chapter_title)}</title>"
            "</head>"
            "<body>"
            f"<h1>{escape(chapter_title)}</h1>"
            f"{paragraph_html}"
            "</body>"
            "</html>"
        )

        epub_book.add_item(item)

        chapter_items.append(
            {
                "title": chapter_title,
                "file_name": item.file_name,
            }
        )

        spine.append(item)

    epub_book.toc = [
        epub.Link(
            chapter["file_name"],
            chapter["title"],
            f"chapter_{index}",
        )
        for index, chapter in enumerate(
            chapter_items,
            start=1,
        )
    ]

    epub_book.spine = spine
    epub_book.add_item(epub.EpubNcx())
    epub_book.add_item(epub.EpubNav())

    epub.write_epub(
        str(output_path),
        epub_book,
    )


def _export_pdf(
    output_path: Path,
    book: dict,
    chapters: list[dict],
) -> None:
    """Create a printable PDF manuscript."""
    try:
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import (
            getSampleStyleSheet,
            ParagraphStyle,
        )
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )

    except ImportError as exc:
        raise ExportError(
            "PDF export requires ReportLab. "
            "Install it using requirements.txt."
        ) from exc

    from html import escape

    document = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=0.8 * inch,
        leftMargin=0.8 * inch,
        topMargin=0.8 * inch,
        bottomMargin=0.8 * inch,
        title=book["title"],
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "BookTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontName="Times-Bold",
        fontSize=24,
        leading=30,
        spaceAfter=24,
    )

    chapter_style = ParagraphStyle(
        "ChapterHeading",
        parent=styles["Heading1"],
        alignment=TA_CENTER,
        fontName="Times-Bold",
        fontSize=18,
        leading=23,
        spaceAfter=24,
    )

    body_style = ParagraphStyle(
        "BookBody",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=11,
        leading=16,
        spaceAfter=10,
        alignment=4,
    )

    story = []

    story.append(
        Paragraph(
            escape(book["title"]),
            title_style,
        )
    )

    story.append(
        Paragraph(
            f"Genre: {escape(book['genre'])}",
            styles["Normal"],
        )
    )

    story.append(
        Paragraph(
            f"Audience: {escape(book['audience'])}",
            styles["Normal"],
        )
    )

    if book.get("premise"):
        story.append(Spacer(1, 24))
        story.append(
            Paragraph(
                "Book Description",
                styles["Heading1"],
            )
        )

        story.append(
            Paragraph(
                escape(book["premise"]),
                body_style,
            )
        )

    for chapter in chapters:
        story.append(PageBreak())

        story.append(
            Paragraph(
                escape(
                    chapter["title"].strip()
                    or f"Chapter {chapter['chapter_number']}"
                ),
                chapter_style,
            )
        )

        for paragraph_text in chapter["content"].split("\n"):
            paragraph_text = paragraph_text.strip()

            if paragraph_text:
                story.append(
                    Paragraph(
                        escape(paragraph_text),
                        body_style,
                    )
                )

    document.build(story)


def export_book(
    book_id: int,
    user_id: int,
    file_format: str,
) -> Path:
    """
    Export a saved book and return the output path.

    Example:
        path = export_book(1, 12345, "docx")
    """
    file_format = file_format.lower().strip()

    if file_format.startswith("."):
        file_format = file_format[1:]

    if file_format not in SUPPORTED_FORMATS:
        raise ExportError(
            "Unsupported format. Choose TXT, DOCX, EPUB, or PDF."
        )

    book, chapters = _get_manuscript(
        book_id,
        user_id,
    )

    config.OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    filename = (
        f"{_safe_filename(book['title'])}"
        f"_book_{book_id}.{file_format}"
    )

    output_path = config.OUTPUT_DIR / filename

    exporters = {
        "txt": _export_txt,
        "docx": _export_docx,
        "epub": _export_epub,
        "pdf": _export_pdf,
    }

    try:
        exporters[file_format](
            output_path,
            book,
            chapters,
        )

    except ExportError:
        raise

    except Exception as exc:
        raise ExportError(
            f"Failed to export the manuscript as "
            f"{file_format.upper()}: {exc}"
        ) from exc

    if not output_path.exists() or output_path.stat().st_size == 0:
        raise ExportError(
            "The export operation completed without producing "
            "a valid output file."
        )

    return output_path
